#!/usr/bin/env python3
"""Update ISRA Column M for production on Hostinger (Node.js app in hPanel).

Replaces Vercel-centric hosting, deployment, access, logging, and rollback text
with Hostinger manual zip deploy, hPanel env vars, and GitHub-based change control.

Usage:
  python3 scripts/apply-isra-hostinger-production.py
"""

from pathlib import Path

import openpyxl

PROJECT_DIR = Path(__file__).resolve().parent.parent
FILE_PATH = PROJECT_DIR / "docs" / "IS Risk Assessment_DC Auto_FILLED.xlsx"
SHEET_NAME = "Analysis & Evaluation"
COLUMN_M = 13

# Row numbers align with control rows in "Analysis & Evaluation" (Column A = control ID).
UPDATES: dict[int, str] = {
    14: (
        "The HALLU platform (Next.js frontend and API route handlers) runs as a single Node.js application "
        "hosted on Hostinger (managed Node.js app in hPanel). "
        "Deployment model: IaaS/PaaS hybrid — application compute on Hostinger VPS/cloud (Singapore data center), "
        "with TLS terminated at Hostinger and the app served via `next start` after `npm run build`. "
        "Production releases are delivered by uploading a source zip built locally (`npm run build:hostinger-zip`); "
        "secrets and environment-specific configuration are not bundled in the zip and are set only in the Hostinger UI. "
        "Third-party SaaS: GitHub (source control), Allianz APIs, and payment providers (Stripe and/or SenangPay) per configuration."
    ),
    16: (
        "Production (Hostinger): Application hosted on Hostinger infrastructure with primary compute in Singapore. "
        "Customer-facing traffic is served over HTTPS to the production domain (hallu.com.my). "
        "Hostinger provides platform-managed SSL, application process management for the Node.js runtime, and hPanel operational access. "
        "DR / recovery: Application source and configuration templates are version-controlled in GitHub; "
        "rollback is performed by redeploying a prior known-good release zip and matching environment variable set documented for that release. "
        "Customer/policy data of record remains in Allianz systems; the platform does not host a separate customer database."
    ),
    21: (
        "DC Auto maintains an internal access control policy for the HALLU platform:\n"
        "1. API credentials (Allianz OAuth keys, Stripe/SenangPay secrets, callback keys) are stored as Hostinger "
        "environment variables accessible only to the server runtime — never exposed to client-side code.\n"
        "2. Source code repository access is restricted to authorized developers via GitHub with branch protection rules.\n"
        "3. Production hosting (Hostinger hPanel → Node.js app) access is limited to authorized administrators with strong authentication.\n"
        "4. No customer-facing admin panel exists — all integration administration is server-side and platform-restricted."
    ),
    22: (
        "User registration and de-registration for the HALLU platform:\n"
        "1. Registration: New developer access is granted by the project owner (CTO) upon formal request. "
        "Access to GitHub, Hostinger hPanel (Node.js app and environment variables), and related tooling requires CTO approval.\n"
        "2. De-registration: Upon staff departure, access is revoked within 24 hours across all platforms "
        "(GitHub, Hostinger hPanel, and any shared deployment credentials).\n"
        "3. Allianz API credentials: Managed by Allianz — DC Auto receives OAuth consumer keys/secrets and rotates per Allianz instruction.\n"
        "4. Payment provider credentials: Issued per provider console; stored only in Hostinger environment variables."
    ),
    23: (
        "Privileged access controls for the HALLU platform:\n"
        "1. Hostinger hPanel (Node.js app, deployments, environment variables): Access restricted to project owner / authorized administrators only.\n"
        "2. GitHub repository: Branch protection on main — pull request review required before merge; direct push to production branch is disabled.\n"
        "3. Environment variables (API keys): Configured only in Hostinger UI; not committed to the repository or included in upload zips.\n"
        "4. Server/runtime: No public administrative interfaces; application exposes only required HTTPS web and API endpoints.\n"
        "5. Policy PDF storage (if enabled): Filesystem path on the VPS configured via environment variable; not web-accessible without authorized tokens."
    ),
    24: (
        "Access rights are reviewed by the project owner (CTO) on a quarterly basis. Given the small team size (2–3 developers), the review involves:\n"
        "1. GitHub repository collaborator list — verify all members are active staff.\n"
        "2. Hostinger hPanel users with access to the Node.js application and environment variables — confirm access is current.\n"
        "3. Payment provider and Allianz integration credentials — verify custodians and rotation schedule.\n"
        "Any discrepancies are resolved immediately upon discovery."
    ),
    25: (
        "Access removal process:\n"
        "1. Upon resignation/termination notification, project owner initiates access revocation within 24 hours.\n"
        "2. Checklist: (a) Remove from GitHub organization/repository, (b) Remove or disable Hostinger hPanel access, "
        "(c) Rotate any shared credentials the departing staff had access to (Hostinger env vars, payment keys, callback secrets as applicable).\n"
        "3. Allianz API credentials: If a key rotation is warranted, DC Auto requests new credentials from Allianz.\n"
        "4. Confirmation of complete access removal is documented via email to management."
    ),
    27: (
        "Information access is restricted as follows:\n"
        "1. Allianz API access: Controlled via OAuth2 client credentials (consumer key + secret). "
        "Only the server-side runtime can authenticate — credentials are never exposed to the browser.\n"
        "2. Application source code: Restricted to authorized developers via GitHub private repository.\n"
        "3. Production environment variables: Accessible only via Hostinger hPanel (authorized administrators).\n"
        "4. Customer data: Temporarily held in browser sessionStorage (client-side isolation per browser tab). "
        "No server-side persistent storage of customer PII beyond operational logs and provider integrations.\n"
        "5. Payment credentials (Stripe secret key, SenangPay HMAC secret): Server-side only in Hostinger environment variables."
    ),
    28: (
        "Secure log-on procedures for system administration:\n"
        "1. GitHub: Email + password with mandatory 2FA (TOTP or security key). Session timeout after inactivity.\n"
        "2. Hostinger hPanel: Strong password and platform-supported MFA where enabled; access limited to authorized personnel.\n"
        "3. Optional VPS/SSH (if used for break-glass maintenance): SSH key-based authentication only; IP restriction where practicable.\n"
        "4. Error messages on login failure are generic where applicable — do not reveal whether username or password is incorrect.\n"
        "5. No customer-facing login system exists — the platform is a public quotation and purchase flow without end-user accounts."
    ),
    29: (
        "Password management:\n"
        "1. Developer access (GitHub, Hostinger hPanel): Minimum 12 characters, complexity enforced by platform where available; 2FA mandatory on GitHub.\n"
        "2. Server access: SSH key-based only when SSH is used — no password-only shell access.\n"
        "3. API credentials (Allianz OAuth, Stripe, SenangPay, callback keys): Machine-generated secrets stored as Hostinger environment variables; not human-memorized passwords.\n"
        "4. No customer passwords exist — the platform does not have user accounts or login functionality.\n"
        "5. Allianz manages their own credential rotation schedule; DC Auto complies with credential updates as issued."
    ),
    30: (
        "Privileged utility programs:\n"
        "1. Production deployment: Controlled release process — local zip build (`npm run build:hostinger-zip`), upload via Hostinger Node.js app, "
        "then `npm install` and `npm run build` on the server per hPanel build settings; process start via `npm start`.\n"
        "2. npm/node package management is controlled via package-lock.json (pinned versions). Only authorized developers can modify dependencies.\n"
        "3. No customer-facing database admin tools — the platform does not operate a separate application database for customer records.\n"
        "4. Diagnostic and admin endpoints are not exposed to the public internet beyond required application APIs.\n"
        "5. Hostinger hPanel deployment and environment configuration changes are restricted to authorized administrators."
    ),
    38: (
        "Change management process:\n"
        "1. All code changes are made via Git feature branches.\n"
        "2. Pull requests require review and approval before merge to main branch.\n"
        "3. Testing: Local development and/or UAT credentials (Allianz UAT, payment test mode) before production release.\n"
        "4. Production deployment: Merge to main, build upload zip locally, deploy to Hostinger Node.js app, verify health on production URL.\n"
        "5. Rollback: Redeploy previous known-good zip from Git tag/commit; restore prior environment variable documentation if changed.\n"
        "6. Environment variable changes require authorized administrator approval and are applied only in Hostinger hPanel (not via source zip).\n"
        "7. API integration changes are tested against Allianz UAT environment before production."
    ),
    39: (
        "Capacity management:\n"
        "1. Hostinger Node.js hosting: Single production Node process (`next start`) sized per Hostinger plan; traffic monitored via Hostinger metrics and application health checks.\n"
        "2. Allianz API: Rate-limited by Allianz per their API gateway policies. The application implements request rate limiting in Next.js middleware to protect upstream APIs.\n"
        "3. Current capacity is adequate for projected transaction volumes as a new digital channel.\n"
        "4. Capacity is reviewed periodically (at least monthly) using Hostinger resource usage and application error/latency trends.\n"
        "5. Scale-up path: Upgrade Hostinger plan or optimize application/build if sustained load requires additional CPU/RAM."
    ),
    40: (
        "Three segregated environments:\n"
        "1. Development: Local developer machines running Next.js dev server (localhost). Uses Allianz UAT credentials and payment test/sandbox configuration.\n"
        "2. Testing/UAT: Same codebase tested against Allianz UAT and payment provider test mode before each production release (no automatic preview hosting).\n"
        "3. Production: Hostinger Node.js deployment (hallu.com.my). Uses Allianz production API and live payment credentials per approved configuration.\n\n"
        "Isolation: Each environment uses separate environment variables (local `.env.local` vs Hostinger UI for production). "
        "No transitive routing between UAT and production credentials."
    ),
    41: (
        "Controls against malware:\n"
        "1. Dependency scanning: npm audit run regularly to detect vulnerable packages. Dependabot alerts enabled on GitHub.\n"
        "2. No arbitrary file uploads: The platform does not accept user file uploads, eliminating a common malware vector.\n"
        "3. Input validation: User inputs validated via Zod schemas — reduces injection of malicious payloads.\n"
        "4. Security headers: Middleware sets restrictive HTTP headers (CSP-related and hardening headers as configured) to mitigate XSS and clickjacking.\n"
        "5. Hosting platform: Hostinger-managed OS patching and platform security for the Node.js runtime environment."
    ),
    42: (
        "Information backup:\n"
        "1. Source code: Fully version-controlled in Git (GitHub). Complete history maintained; repository backed up by GitHub infrastructure.\n"
        "2. Configuration: Production environment variables documented in secure internal records (values stored only in Hostinger UI).\n"
        "3. Customer data: No persistent customer/policy database on the application server — policy data of record resides in Allianz systems.\n"
        "4. Application releases: Each production deploy is tied to a Git commit; prior releases can be rebuilt and redeployed from Git history.\n"
        "5. Policy PDF files (if stored on VPS): Filesystem backups are the operator's responsibility per Hostinger backup options and internal runbooks."
    ),
    43: (
        "Event logging:\n"
        "1. Application logs: API interactions with Allianz and payment providers logged with timestamps, status codes, and error context on the Node.js server.\n"
        "2. Hostinger logs: Application stdout/stderr and deployment/build logs available via hPanel for the Node.js app.\n"
        "3. Deployment audit: Git commit history and pull request records provide change traceability; Hostinger records deployment actions by authorized users.\n"
        "4. Logs are not editable by end users; production log access is limited to authorized administrators.\n"
        "5. Log review: Logs are reviewed upon incident or anomaly detection; periodic review during production operations.\n"
        "6. Retention follows Hostinger platform limits; critical incidents may export logs for longer retention."
    ),
    45: (
        "Cryptographic controls:\n"
        "1. Data in transit: All communications use TLS 1.2+ (HTTPS enforced). Hostinger provides SSL/TLS for the production domain.\n"
        "2. API authentication: OAuth2 with client credentials grant to Allianz; tokens are short-lived.\n"
        "3. Payment security: Stripe Checkout / SenangPay integration uses provider-defined signing (e.g. HMAC-SHA256 for SenangPay). "
        "Secret keys stored in Hostinger environment variables only.\n"
        "4. Allianz callback verification: Shared secret / API key validation on incoming webhooks where configured.\n"
        "5. Policy document access: Short-lived signed tokens for authorized download paths where implemented."
    ),
    53: (
        "The HALLU platform is a Progressive Web Application (PWA) / responsive website — not a native mobile app.\n"
        "1. It runs in the customer's mobile browser (Safari/Chrome) — no app store download required.\n"
        "2. The web application is served over HTTPS from the production Hostinger deployment.\n"
        "3. Browser security model provides sandboxing — the web app cannot access device storage, contacts, or other apps beyond normal web APIs.\n"
        "4. Content is served from the currently deployed production build on Hostinger.\n"
        "5. Mobile OS security updates are the responsibility of the device owner; the application does not require a separate mobile OS patch channel."
    ),
    62: (
        "Cyber security controls:\n"
        "1. DDoS / availability: Hostinger platform networking and TLS edge; application-level rate limiting in Next.js middleware on sensitive routes.\n"
        "2. Input validation: Zod schema validation helps prevent injection attacks (XSS, malformed payloads).\n"
        "3. Security headers: HTTP security headers applied via application middleware.\n"
        "4. Secrets: No secrets in client bundles; keys only in Hostinger environment variables.\n"
        "5. Dependency hygiene: Locked dependencies and periodic vulnerability review (npm audit / Dependabot).\n"
        "6. Payment flows: Card data handled by Stripe hosted checkout where used; SenangPay redirect flow per provider security model."
    ),
    67: (
        "Change control procedures:\n"
        "1. Developer creates feature branch from main.\n"
        "2. Code changes are committed with descriptive messages.\n"
        "3. Pull request created with description of changes and testing performed.\n"
        "4. Peer review required — at least one approval before merge.\n"
        "5. After approval, merge to main; authorized administrator builds upload zip and deploys to Hostinger Node.js app.\n"
        "6. Post-deploy verification on production URL and smoke test of quote/payment critical paths.\n"
        "7. Rollback: Redeploy previous known-good zip and env configuration from documented release record.\n"
        "8. Emergency changes: Same PR/review where time permits; otherwise deploy with retrospective documentation within 24 hours."
    ),
    76: (
        "Response to information security incidents:\n"
        "1. Immediate containment: Rotate compromised secrets in Hostinger environment variables, disable affected routes, or take the Node.js app offline via Hostinger if necessary.\n"
        "2. Investigation: Review Hostinger application logs, recent GitHub commits/deployments, and integration provider dashboards to determine scope and root cause.\n"
        "3. Remediation: Fix vulnerability, rotate credentials, redeploy patched build to Hostinger.\n"
        "4. Communication: Notify Allianz of incidents affecting the integration within 4 hours.\n"
        "5. Post-incident: Document lessons learned, update security controls, conduct review.\n"
        "6. Credential rotation: Allianz API keys, payment provider secrets, and callback keys rotated if any compromise is suspected."
    ),
}


def main() -> None:
    if not FILE_PATH.is_file():
        raise SystemExit(f"ISRA file not found: {FILE_PATH}")

    wb = openpyxl.load_workbook(FILE_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {SHEET_NAME}")

    ws = wb[SHEET_NAME]

    for row_num, response in sorted(UPDATES.items()):
        ctrl_id = ws.cell(row=row_num, column=1).value
        key_area = ws.cell(row=row_num, column=2).value
        ws.cell(row=row_num, column=COLUMN_M).value = response
        print(f"  ✓ Row {row_num} ({ctrl_id} — {key_area})")

    wb.save(FILE_PATH)
    print(f"\n✅ ISRA Column M updated for Hostinger production: {FILE_PATH}")
    print(f"   Rows updated: {len(UPDATES)}")


if __name__ == "__main__":
    main()
